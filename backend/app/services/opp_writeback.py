"""Write inline opportunity edits straight into the ``Opp. FYxx`` sheets.

The row is located again at write time (by Opp ID MMS, falling back to the
row recorded at load time when the project name still matches), so edits land
on the right row even if the file changed since it was loaded. Only the input
sheets are ever written; formula sheets recompute in Excel.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl

from app.models import Opportunity
from app.services import xlsx_patch
from app.services.excel_reader import norm_opp_id

# Opportunity attribute -> column header in the Opp sheets (lower-case).
EDITABLE_COLUMNS = {
    "mms_status": "mms status",
    "acn_tool_status": "stato acn tool",
}


def _locate(path: Path, opp: Opportunity) -> tuple[int, dict[str, int], list]:
    """Return (excel_row, {header: col_index0}, row_values) of ``opp`` in its sheet."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if opp.source_sheet not in wb.sheetnames:
            raise LookupError(f"Foglio '{opp.source_sheet}' non trovato: ricarica i dati")
        rows = [list(r) for r in wb[opp.source_sheet].iter_rows(min_row=1, min_col=1, values_only=True)]
    finally:
        wb.close()
    hdr = next((i for i, r in enumerate(rows[:5])
                if any(isinstance(v, str) and v.strip() == "Contract" for v in r)), None)
    if hdr is None:
        raise LookupError("Intestazione del foglio opportunità non trovata")
    cols = {str(v).strip().lower(): i for i, v in enumerate(rows[hdr]) if v is not None}
    id_col, name_col = cols.get("opp id mms"), cols.get("project")

    def cell(r, c):
        return r[c] if c is not None and c < len(r) else None

    target = norm_opp_id(opp.opp_id_mms)
    if target:
        hits = [i + 1 for i, r in enumerate(rows) if i > hdr and norm_opp_id(cell(r, id_col)) == target]
        if len(hits) == 1:
            return hits[0], cols, rows[hits[0] - 1]
        if opp.source_row in hits:
            return opp.source_row, cols, rows[opp.source_row - 1]
    if opp.source_row and opp.source_row <= len(rows):
        r = rows[opp.source_row - 1]
        if str(cell(r, name_col) or "").strip() == (opp.name or "").strip():
            return opp.source_row, cols, r
    raise LookupError("Riga dell'opportunità non trovata nel file Excel: ricarica i dati")


def _coerce(existing, value):
    """Keep the cell type: '1' stays numeric if the cell held a number."""
    if value is None or value == "":
        return None
    if isinstance(existing, (int, float)) and str(value).strip().lstrip("-").isdigit():
        return int(str(value).strip())
    return str(value)


def write_opportunity_fields(path: Path, opp: Opportunity, changes: dict) -> Path:
    """Write the editable ``changes`` of ``opp`` to Excel. Returns the backup path."""
    try:
        row, cols, values = _locate(path, opp)
    except PermissionError as exc:
        raise xlsx_patch.WorkbookLockedError(
            "Il file Excel è aperto o bloccato: chiudilo in Excel e riprova."
        ) from exc
    cells = {}
    for attr, value in changes.items():
        header = EDITABLE_COLUMNS.get(attr)
        if header is None or header not in cols:
            continue
        ci = cols[header]
        existing = values[ci] if ci < len(values) else None
        cells[f"{xlsx_patch.col_letter(ci + 1)}{row}"] = _coerce(existing, value)
    if not cells:
        raise LookupError("Nessuna colonna modificabile trovata nel foglio")
    return xlsx_patch.set_cells(path, opp.source_sheet, cells)
