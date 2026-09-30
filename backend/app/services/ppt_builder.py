"""Build the executive account-summary deck in a chosen template.

The same five slides (cover, KPI, executive summary, economics, pipeline) are
produced for every template:

* built-in style -> blank 16:9 deck, backgrounds and accents drawn with the
  style's palette and fonts;
* uploaded template -> the uploaded .pptx itself: example slides removed, new
  slides created from its own layouts (cover layout for the cover, title-only /
  title+content layouts for the others), titles written into the layout's
  placeholders, cards and charts coloured with the template's theme.
"""
from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from app.services.ppt_templates import BUILTIN, TEMPLATE_DIR, Style, get_template, style_from_analysis

WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def _eur(v: float) -> str:
    return f"€ {v:,.0f}".replace(",", ".")


class Deck:
    """Thin wrapper hiding the differences between built-in and uploaded templates."""

    def __init__(self, template_id: str):
        self.uploaded = template_id not in BUILTIN
        if self.uploaded:
            meta = get_template(template_id)
            if meta is None:
                raise ValueError("Template non trovato")
            self.style = style_from_analysis(meta["name"], meta["analysis"])
            self.prs = Presentation(str(TEMPLATE_DIR / template_id / "template.pptx"))
            self._clear_slides()
            self._pick_layouts()
        else:
            self.style = BUILTIN[template_id]["style"]
            self.prs = Presentation()
            self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
        self.W, self.H = self.prs.slide_width, self.prs.slide_height

    # ---------------- uploaded-template plumbing ----------------
    def _clear_slides(self) -> None:
        """Drop the example slides, keeping masters/layouts/theme."""
        sld_ids = self.prs.slides._sldIdLst  # noqa: SLF001 - no public API for removal
        for sld_id in list(sld_ids):
            self.prs.part.drop_rel(sld_id.rId)
            sld_ids.remove(sld_id)

    @staticmethod
    def _ph_types(layout) -> set:
        types = set()
        for ph in layout.placeholders:
            try:
                types.add(ph.placeholder_format.type)
            except (AttributeError, ValueError):
                pass
        return types

    def _pick_layouts(self) -> None:
        layouts = list(self.prs.slide_layouts)
        title_types = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
        cover = next((l for l in layouts if PP_PLACEHOLDER.CENTER_TITLE in self._ph_types(l)), None)
        title_only = next((l for l in layouts if self._ph_types(l) & title_types
                           and not self._ph_types(l) & {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT,
                                                        PP_PLACEHOLDER.SUBTITLE}), None)
        content = next((l for l in layouts if PP_PLACEHOLDER.TITLE in self._ph_types(l)
                        and self._ph_types(l) & {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT}), None)
        self.cover_layout = cover or layouts[0]
        self.title_layout = title_only or content or layouts[min(1, len(layouts) - 1)]

    @staticmethod
    def _drop_empty_placeholders(slide) -> None:
        for ph in list(slide.placeholders):
            if ph.has_text_frame and not ph.text_frame.text.strip():
                ph._element.getparent().remove(ph._element)  # noqa: SLF001

    # ---------------- drawing helpers ----------------
    def rect(self, slide, x, y, w, h, color: RGBColor, shape=MSO_SHAPE.RECTANGLE):
        sp = slide.shapes.add_shape(shape, x, y, w, h)
        sp.fill.solid()
        sp.fill.fore_color.rgb = color
        sp.line.fill.background()
        sp.shadow.inherit = False
        return sp

    def text(self, slide, x, y, w, h, text, size, color, bold=False, heading=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        p = tf.paragraphs[0]
        p.alignment = align
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = self.style.heading_font if heading else self.style.body_font
        run.font.color.rgb = color
        return tb

    def bullets(self, slide, x, y, w, h, title, items, accent: RGBColor):
        s = self.style
        self.rect(slide, x, y, Inches(0.1), h, accent)
        self.text(slide, x + Inches(0.3), y - Inches(0.05), w, Inches(0.5), title, 18,
                  s.rgb("ink"), bold=True, heading=True)
        tb = slide.shapes.add_textbox(x + Inches(0.3), y + Inches(0.5), w - Inches(0.4), h - Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, it in enumerate(items or ["—"]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(8)
            run = p.add_run()
            run.text = "•  " + str(it)
            run.font.size = Pt(14)
            run.font.name = s.body_font
            run.font.color.rgb = s.rgb("ink")

    def content_slide(self, title: str):
        """New content slide with its title; returns (slide, top of the free area)."""
        s = self.style
        if self.uploaded:
            slide = self.prs.slides.add_slide(self.title_layout)
            if slide.shapes.title is not None:
                slide.shapes.title.text = title
                top = slide.shapes.title.top + slide.shapes.title.height + Inches(0.2)
            else:
                self.text(slide, Inches(0.6), Inches(0.3), self.W - Inches(1.2), Inches(0.7),
                          title, 24, s.rgb("ink"), bold=True, heading=True)
                top = Inches(1.2)
            self._drop_empty_placeholders(slide)
            return slide, top
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        if s.dark_mode:
            self.rect(slide, 0, 0, self.W, self.H, s.rgb("background"))
            self.rect(slide, Inches(0.6), Inches(1.05), Inches(1.2), Inches(0.06), s.rgb("primary"))
            self.text(slide, Inches(0.6), Inches(0.3), self.W - Inches(1.2), Inches(0.7),
                      title, 26, s.rgb("ink"), bold=True, heading=True)
        elif s.cover_dark:
            self.rect(slide, 0, 0, self.W, Inches(1.0), s.rgb("primary"))
            self.text(slide, Inches(0.6), Inches(0.24), self.W - Inches(1.2), Inches(0.6),
                      title, 24, WHITE, bold=True, heading=True)
        else:
            self.text(slide, Inches(0.6), Inches(0.35), self.W - Inches(1.2), Inches(0.7),
                      title, 28, s.rgb("dark"), heading=True)
            self.rect(slide, Inches(0.6), Inches(1.05), self.W - Inches(1.2), Emu(12700), s.rgb("primary"))
        return slide, Inches(1.4)

    def color_chart(self, chart, colors: list[RGBColor]):
        for i, series in enumerate(chart.plots[0].series):
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = colors[i % len(colors)]
        chart.value_axis.tick_labels.number_format = '#,##0'
        chart.value_axis.tick_labels.number_format_is_linked = False
        for axis in (chart.category_axis, chart.value_axis):
            axis.tick_labels.font.size = Pt(11)
            axis.tick_labels.font.name = self.style.body_font
            if self.style.dark_mode:
                axis.tick_labels.font.color.rgb = self.style.rgb("ink")

    # ---------------- slides ----------------
    def cover(self, title: str, subtitle: str):
        s = self.style
        if self.uploaded:
            slide = self.prs.slides.add_slide(self.cover_layout)
            if slide.shapes.title is not None:
                slide.shapes.title.text = title
            sub = next((ph for ph in slide.placeholders
                        if ph.placeholder_format.type in (PP_PLACEHOLDER.SUBTITLE, PP_PLACEHOLDER.BODY)), None)
            if sub is not None:
                sub.text = subtitle
            self._drop_empty_placeholders(slide)
            return
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        if s.cover_dark or s.dark_mode:
            self.rect(slide, 0, 0, self.W, self.H, s.rgb("dark"))
            self.rect(slide, 0, Inches(5.1), self.W, Inches(0.12), s.rgb("primary"))
            fg, sub_fg = WHITE, s.rgb("light") if not s.dark_mode else s.rgb("muted")
        else:
            self.rect(slide, 0, 0, Inches(0.35), self.H, s.rgb("primary"))
            fg, sub_fg = s.rgb("dark"), s.rgb("muted")
        self.text(slide, Inches(0.9), Inches(2.4), Inches(11.5), Inches(1.6), title, 40, fg, bold=True, heading=True)
        self.text(slide, Inches(0.95), Inches(4.1), Inches(11), Inches(0.9), subtitle, 18, sub_fg)
        self.text(slide, Inches(0.95), Inches(6.7), Inches(11), Inches(0.5),
                  "PM Control Center · Account Review", 12, s.rgb("muted"))

    def kpis(self, data: dict):
        s = self.style
        slide, top = self.content_slide("KPI dell'account")
        # On dark slides "dark" would vanish into the background: use "light" instead.
        second = s.rgb("light") if s.dark_mode else s.rgb("dark")
        ci_color = s.rgb("good") if data["ci_pct"] >= 0.35 else s.rgb("warn") if data["ci_pct"] >= 0.30 else s.rgb("bad")
        cards = [
            ("Ricavi", _eur(data["revenues"]), s.rgb("primary")),
            ("Costi totali", _eur(data["costs"]), second),
            ("Contribution Income", _eur(data["ci"]), s.rgb("primary")),
            ("CCI %", f"{data['ci_pct']*100:.1f}%", ci_color),
            ("Contratti", str(data["contracts"]), second),
            ("Risorse", str(data["resources"]), s.rgb("primary")),
            ("Vinto (CloseWon)", _eur(data["won"]), s.rgb("good")),
            ("Pipeline aperta", _eur(data["pipeline_total"]), s.rgb("warn")),
        ]
        cols, gap = 4, Inches(0.3)
        margin = Inches(0.6)
        cw = int((self.W - 2 * margin - gap * (cols - 1)) / cols)
        avail_h = self.H - top - Inches(0.5)
        ch = int(min(Inches(1.9), (avail_h - Inches(0.35)) / 2))
        value_pt = 24 if cw >= Inches(2.6) else 18  # narrow cards (4:3 templates)
        for i, (label, value, color) in enumerate(cards):
            r, c = divmod(i, cols)
            card = self.rect(slide, margin + c * (cw + gap), top + r * (ch + Inches(0.35)), cw, ch,
                             color, MSO_SHAPE.ROUNDED_RECTANGLE)
            tf = card.text_frame
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run(); run.text = value
            run.font.size = Pt(value_pt); run.font.bold = True; run.font.color.rgb = WHITE
            run.font.name = s.heading_font
            p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
            run2 = p2.add_run(); run2.text = label
            run2.font.size = Pt(11); run2.font.color.rgb = WHITE; run2.font.name = s.body_font

    def summary(self, nar: dict):
        s = self.style
        slide, top = self.content_slide("Executive Summary")
        h = self.H - top - Inches(0.5)
        left_w = int((self.W - Inches(1.5)) * 0.6)
        right_x = Inches(0.6) + left_w + Inches(0.3)
        right_w = self.W - right_x - Inches(0.6)
        self.bullets(slide, Inches(0.6), top, left_w, h, "Sintesi",
                     nar["executive_summary"] or ["Riepilogo non disponibile."], s.rgb("primary"))
        half = int((h - Inches(0.2)) / 2)
        self.bullets(slide, right_x, top, right_w, half, "Punti di forza", nar["highlights"], s.rgb("good"))
        self.bullets(slide, right_x, top + half + Inches(0.2), right_w, half, "Raccomandazioni",
                     nar["recommendations"], s.rgb("warn"))

    def economics(self, data: dict):
        s = self.style
        slide, top = self.content_slide("Andamento economico")
        chart_data = CategoryChartData()
        chart_data.categories = ["Ricavi", "Costi", "CI"]
        chart_data.add_series("EUR", (data["revenues"], data["costs"], data["ci"]))
        w = int((self.W - Inches(1.2)) * 0.5)
        gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), top, w,
                                    self.H - top - Inches(0.5), chart_data)
        gf.chart.has_legend = False
        self.color_chart(gf.chart, [s.rgb("primary")])
        x = Inches(0.6) + w + Inches(0.5)
        tw = self.W - x - Inches(0.6)
        self.text(slide, x, top + Inches(0.3), tw, Inches(0.6), f"CCI = {data['ci_pct']*100:.1f}%",
                  24, s.rgb("ink"), bold=True, heading=True)
        target = "≥ 35% (in target)" if data["ci_pct"] >= 0.35 else "sotto il target del 35%"
        self.text(slide, x, top + Inches(1.0), tw, Inches(0.6), target, 14, s.rgb("muted"))
        self.text(slide, x, top + Inches(1.8), tw, Inches(2.0),
                  f"Ricavi {_eur(data['revenues'])}\nCosti {_eur(data['costs'])}\nCI {_eur(data['ci'])}",
                  14, s.rgb("ink"))

    def pipeline(self, data: dict):
        if not data["pipeline_by_stage"]:
            return
        s = self.style
        slide, top = self.content_slide("Pipeline per stage")
        chart_data = CategoryChartData()
        items = list(data["pipeline_by_stage"].items())
        chart_data.categories = [k for k, _ in items]
        chart_data.add_series("Valore", tuple(v for _, v in items))
        gf = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.6), top, self.W - Inches(1.2),
                                    self.H - top - Inches(0.5), chart_data)
        gf.chart.has_legend = False  # single series: the axis already names the stages
        self.color_chart(gf.chart, [s.rgb("primary")])

    def save(self) -> BytesIO:
        out = BytesIO()
        self.prs.save(out)
        out.seek(0)
        return out


def build_deck(template_id: str, data: dict, nar: dict) -> BytesIO:
    deck = Deck(template_id)
    subtitle = nar.get("subtitle") or f"{', '.join(data['clients']) or 'Account'} · {date.today().strftime('%d/%m/%Y')}"
    deck.cover("Account Summary", subtitle)
    deck.kpis(data)
    deck.summary(nar)
    deck.economics(data)
    deck.pipeline(data)
    return deck.save()
