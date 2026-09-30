"""PowerPoint templates for the executive account summary.

Two kinds of templates:

* **Built-in styles**: palettes/fonts drawn on a blank 16:9 deck.
* **Uploaded templates**: example decks (.pptx) uploaded by the user. Each one is
  analysed (slide size, theme palette, heading/body fonts, layouts, structure
  and content types of its slides) and stored in
  ``%APPDATA%/PMControlCenter/ppt_templates/<id>/``. The deck is then generated
  *on top of the uploaded file*: its example slides are removed and new slides
  are created from its own layouts, so masters, backgrounds, logos and fonts
  are exactly those of the template.
"""
from __future__ import annotations

import json
import re
import shutil
import uuid
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from app.core.config import APP_DATA_DIR

TEMPLATE_DIR = APP_DATA_DIR / "ppt_templates"
_EMU_PER_INCH = 914400


@dataclass
class Style:
    """Colours and fonts used to draw the deck."""
    name: str
    primary: str
    dark: str
    ink: str
    muted: str
    light: str
    background: str
    good: str = "10B981"
    warn: str = "F59E0B"
    bad: str = "EF4444"
    heading_font: str = "Calibri"
    body_font: str = "Calibri"
    cover_dark: bool = True  # dark cover with light text
    dark_mode: bool = False  # dark content slides
    extra: dict = field(default_factory=dict)

    def rgb(self, attr: str) -> RGBColor:
        return RGBColor.from_string(getattr(self, attr))


BUILTIN: dict[str, dict] = {
    "builtin:accenture": {
        "name": "Accenture viola",
        "description": "Copertina viola scura, accenti viola, sfondo bianco (stile attuale).",
        "style": Style("Accenture viola", primary="7C3AED", dark="4C1D95", ink="1E293B",
                       muted="64748B", light="F1F5F9", background="FFFFFF",
                       heading_font="Arial", body_font="Arial"),
    },
    "builtin:corporate": {
        "name": "Corporate chiaro",
        "description": "Minimal su bianco, blu istituzionale, titoli sottili, adatto a stampa.",
        "style": Style("Corporate chiaro", primary="1D4ED8", dark="1E3A8A", ink="111827",
                       muted="6B7280", light="EFF6FF", background="FFFFFF",
                       heading_font="Segoe UI Light", body_font="Segoe UI", cover_dark=False),
    },
    "builtin:dark": {
        "name": "Dark executive",
        "description": "Slide scure ad alto contrasto con accenti ciano, per proiezione.",
        "style": Style("Dark executive", primary="06B6D4", dark="0B1220", ink="E5E7EB",
                       muted="94A3B8", light="1F2937", background="111827",
                       heading_font="Segoe UI Semibold", body_font="Segoe UI", dark_mode=True),
    },
}


# --------------------------------------------------------------------------- #
# analysis of an uploaded deck
# --------------------------------------------------------------------------- #
def _theme_xml(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        themes = sorted(n for n in z.namelist() if re.match(r"ppt/theme/theme\d+\.xml$", n))
        return z.read(themes[0]).decode("utf-8") if themes else ""


def _theme_palette(xml: str) -> dict[str, str]:
    palette = {}
    block = re.search(r"<a:clrScheme\b.*?</a:clrScheme>", xml, re.S)
    if not block:
        return palette
    for slot in ("dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3",
                 "accent4", "accent5", "accent6", "hlink"):
        m = re.search(rf"<a:{slot}>(.*?)</a:{slot}>", block.group(0), re.S)
        if not m:
            continue
        c = re.search(r'srgbClr val="([0-9A-Fa-f]{6})"', m.group(1)) or \
            re.search(r'lastClr="([0-9A-Fa-f]{6})"', m.group(1))
        if c:
            palette[slot] = c.group(1).upper()
    return palette


def _theme_fonts(xml: str) -> dict[str, str]:
    out = {}
    for kind, key in (("majorFont", "heading"), ("minorFont", "body")):
        m = re.search(rf"<a:{kind}>.*?<a:latin typeface=\"([^\"]*)\"", xml, re.S)
        if m and m.group(1):
            out[key] = m.group(1)
    return out


def _placeholder_kinds(shapes) -> list[str]:
    kinds = []
    for ph in shapes.placeholders if hasattr(shapes, "placeholders") else []:
        try:
            kinds.append(ph.placeholder_format.type.name if ph.placeholder_format.type else "OBJECT")
        except (AttributeError, ValueError):
            kinds.append("OBJECT")
    return kinds


def _shape_kind(sh) -> str:
    if sh.has_chart if hasattr(sh, "has_chart") else False:
        return "grafico"
    if sh.has_table if hasattr(sh, "has_table") else False:
        return "tabella"
    st = sh.shape_type
    if st == MSO_SHAPE_TYPE.PICTURE:
        return "immagine"
    if st == MSO_SHAPE_TYPE.GROUP:
        return "gruppo"
    if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
        return "testo"
    return "forma"


def analyze_pptx(path: Path) -> dict:
    prs = Presentation(str(path))
    xml = _theme_xml(path)
    palette = _theme_palette(xml)
    fonts = _theme_fonts(xml)
    layouts = []
    for i, layout in enumerate(prs.slide_layouts):
        layouts.append({"index": i, "name": layout.name, "placeholders": _placeholder_kinds(layout)})
    slides = []
    totals: dict[str, int] = {}
    for i, slide in enumerate(prs.slides, start=1):
        kinds: dict[str, int] = {}
        for sh in slide.shapes:
            k = _shape_kind(sh)
            kinds[k] = kinds.get(k, 0) + 1
            totals[k] = totals.get(k, 0) + 1
        title = slide.shapes.title.text_frame.text.strip() if slide.shapes.title is not None else ""
        slides.append({"n": i, "layout": slide.slide_layout.name, "title": title[:80], "content": kinds})
    return {
        "slide_width_in": round(prs.slide_width / _EMU_PER_INCH, 2),
        "slide_height_in": round(prs.slide_height / _EMU_PER_INCH, 2),
        "palette": palette,
        "fonts": fonts,
        "layouts": layouts,
        "slides": slides,
        "content_totals": totals,
    }


def style_from_analysis(name: str, a: dict) -> Style:
    """Derive drawing colours/fonts from a template's theme."""
    p = a.get("palette", {})
    f = a.get("fonts", {})
    return Style(
        name,
        primary=p.get("accent1", "7C3AED"),
        dark=p.get("dk2", p.get("accent2", "1E293B")),
        ink=p.get("dk1", "1E293B"),
        muted=p.get("accent3", "64748B"),
        light=p.get("lt2", "F1F5F9"),
        background=p.get("lt1", "FFFFFF"),
        heading_font=f.get("heading", "Calibri"),
        body_font=f.get("body", "Calibri"),
    )


# --------------------------------------------------------------------------- #
# storage
# --------------------------------------------------------------------------- #
def _meta_path(tid: str) -> Path:
    return TEMPLATE_DIR / tid / "meta.json"


def save_template(data: bytes, filename: str) -> dict:
    if not filename.lower().endswith(".pptx"):
        raise ValueError("Carica un file PowerPoint .pptx")
    tid = uuid.uuid4().hex[:12]
    folder = TEMPLATE_DIR / tid
    folder.mkdir(parents=True, exist_ok=True)
    pptx_path = folder / "template.pptx"
    pptx_path.write_bytes(data)
    try:
        analysis = analyze_pptx(pptx_path)
    except Exception as exc:  # noqa: BLE001 - not a valid deck
        shutil.rmtree(folder, ignore_errors=True)
        raise ValueError(f"File PowerPoint non leggibile: {exc}") from exc
    meta = {"id": tid, "name": Path(filename).stem, "filename": filename,
            "uploaded_at": datetime.now().isoformat(timespec="seconds"), "analysis": analysis}
    _meta_path(tid).write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    return meta


def uploaded_templates() -> list[dict]:
    if not TEMPLATE_DIR.is_dir():
        return []
    out = []
    for meta in sorted(TEMPLATE_DIR.glob("*/meta.json")):
        try:
            out.append(json.loads(meta.read_text(encoding="utf-8")))
        except (OSError, json.JSONDecodeError):
            continue
    return sorted(out, key=lambda m: m.get("uploaded_at", ""), reverse=True)


def get_template(tid: str) -> dict | None:
    if tid in BUILTIN:
        return None
    p = _meta_path(tid)
    return json.loads(p.read_text(encoding="utf-8")) if p.is_file() else None


def delete_template(tid: str) -> bool:
    if not re.fullmatch(r"[0-9a-f]{12}", tid):
        return False
    folder = TEMPLATE_DIR / tid
    if folder.is_dir():
        shutil.rmtree(folder)
        return True
    return False


def list_templates() -> list[dict]:
    items = [{"id": k, "kind": "builtin", "name": v["name"], "description": v["description"],
              "palette": {"primary": v["style"].primary, "dark": v["style"].dark,
                          "background": v["style"].background},
              "fonts": {"heading": v["style"].heading_font, "body": v["style"].body_font}}
             for k, v in BUILTIN.items()]
    for m in uploaded_templates():
        a = m["analysis"]
        items.append({"id": m["id"], "kind": "uploaded", "name": m["name"],
                      "description": f"{len(a['slides'])} slide di esempio · {len(a['layouts'])} layout · "
                                     f"{a['slide_width_in']}×{a['slide_height_in']} in",
                      "palette": a.get("palette", {}), "fonts": a.get("fonts", {}),
                      "analysis": a, "uploaded_at": m.get("uploaded_at")})
    return items
